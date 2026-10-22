"""
thickness_stats.py  --  run with normal Python (NOT inside CFD-VIEW):   python thickness_stats.py

Reads the files written by the CFD-VIEW scripts
    <n>-Hf.txt, <n>-Zr.txt      (per-node thickness, from hzo_halfcycle_thickness.py)
    areal_summary.csv           (optional, from hzo_halfcycle_areal.py)
and produces, in the sub-folder OUT_SUBDIR of DATA_DIR (created automatically)
    thickness_stats.xlsx  Excel workbook with all tables below (one sheet each) + a Summary sheet
    thickness_stats.csv   one row per half cycle  (mean, min, max, std, increment, node variation ...)
    cycle_stats.csv       one row per full cycle  (added Hf, added Zr, GPC, total after cycle, Hf:Zr ratio)
    deposition_error.csv  one row per cycle       (error % of added thickness vs mean and vs mode)
    plots (PNG)
        plot_increment_per_halfcycle.png   thickness added in every half cycle
        plot_added_Hf_Zr_separate.png      added by Hf / by Zr half cycle against cycle number
        plot_growth_per_cycle.png          growth per cycle (Hf + Zr) with mean and mode
        plot_HfZr_ratio.png                Hf:Zr deposition thickness ratio per cycle
        plot_node_variation.png            % variation of node max / min wrt mean in every half-cycle file
        plot_deposition_error.png          error % of added thickness wrt mean and wrt mode

Definitions
    mean thickness            = mean of the thickness over all nodes (after that half cycle)
    half-cycle increment      = mean(this half cycle) - mean(previous half cycle)
                                (for the very first half cycle the previous mean is 0)
    growth per cycle (GPC)    = mean(after Zr of cycle n) - mean(after Zr of cycle n-1)
    total after cycle         = mean total thickness after the last (Zr) half cycle of that cycle
    Hf:Zr ratio               = thickness added by the Hf half cycle / thickness added by the Zr
                                half cycle of the same cycle
    node variation (+max %)   = (max - mean) / mean * 100      over the nodes of one file
    node variation (-min %)   = (min - mean) / mean * 100      (negative = below the mean)
    error vs mean (%)         = (added - mean of added) / mean of added * 100
    error vs mode (%)         = (added - mode of added) / mode of added * 100
                                mean and mode are taken over the cycles after SKIP_CYCLES
    mode                      = most frequent value after rounding to MODE_DECIMALS; if no value
                                repeats, the centre of the most populated histogram bin
Requires: numpy, pandas, matplotlib, openpyxl   (pip install numpy pandas matplotlib openpyxl)
"""
import os
import re
import sys
import time
import numpy as np
import pandas as pd
import matplotlib
if "--noshow" in sys.argv:
    matplotlib.use("Agg")
import matplotlib.pyplot as plt

# ----------------------------- USER SETTINGS -----------------------------

DATA_DIR    = r"C:\Tests\Calibration_w_o_recomb\Stats\0.125_out"  # folder with 1-Hf.txt, 1-Zr.txt, ...
HALF_CYCLES = ["Hf", "Zr"]                       # order inside one cycle
SCALE       = 1e10            # multiply file values by this (e.g. 1e9 for m -> nm)
UNIT        = "A"   # label used on plots, e.g. "nm"

BASIS         = "mean_nodal" # mean used for ratio / GPC / error stats ("mean_nodal" or e.g.
                             # "mean_area_weighted" if areal_summary.csv is present)
SKIP_CYCLES   = 0            # first N cycles left out of the reference mean / mode (e.g. nucleation)
MODE_DECIMALS = 3            # values are rounded to this many decimals (in UNIT) to find the mode
EXCEL_NAME    = "thickness_stats.xlsx"
OUT_SUBDIR    = "thickness_stats_results"   # results go into DATA_DIR\OUT_SUBDIR (created if missing)

FIG_W, FIG_H = 20, 9         # figure size in inches (bigger = less crowded x axis)
MAX_XLABELS  = 30            # at most this many x tick labels are written
Y_PAD        = 0.08          # empty margin above/below the data, as a fraction of its range
# --------------------------------------------------------------------------


if len(sys.argv) > 1 and not sys.argv[1].startswith("--"):
    DATA_DIR = sys.argv[1]

OUT_DIR = os.path.join(DATA_DIR, OUT_SUBDIR)
os.makedirs(OUT_DIR, exist_ok=True)

hc_rank = {h: i for i, h in enumerate(HALF_CYCLES)}
H1, H2 = HALF_CYCLES[0], HALF_CYCLES[-1]
RATIO_COL = "ratio_%s_%s" % (H1, H2)


def list_files():
    out = []
    for f in os.listdir(DATA_DIR):
        m = re.match(r"^(\d+)-([A-Za-z]+)\.txt$", f)
        if m and m.group(2) in hc_rank:
            out.append((int(m.group(1)), m.group(2), os.path.join(DATA_DIR, f)))
    out.sort(key=lambda t: (t[0], hc_rank[t[1]]))
    return out


def read_txt(path):
    df = pd.read_csv(path, sep="\t", comment="#")
    return df


def pct(a, ref):
    a = np.asarray(a, dtype=float)
    ref = np.asarray(ref, dtype=float)
    with np.errstate(divide="ignore", invalid="ignore"):
        return np.where(ref != 0, (a - ref) / ref * 100.0, np.nan)


def mode_of(v):
    """Mode of continuous data: most frequent rounded value, else histogram peak."""
    v = np.asarray(v, dtype=float)
    v = v[np.isfinite(v)]
    if v.size == 0:
        return np.nan, "no data"
    r = np.round(v, MODE_DECIMALS)
    vals, counts = np.unique(r, return_counts=True)
    if counts.max() > 1:
        return float(vals[counts.argmax()]), "most frequent value rounded to %d decimals (%d of %d)" \
            % (MODE_DECIMALS, counts.max(), v.size)
    cnt, edges = np.histogram(v, bins="auto")
    i = int(cnt.argmax())
    return float(0.5 * (edges[i] + edges[i + 1])), "histogram peak (%d bins, no repeated value)" % len(cnt)


# ------------------------- per-node statistics ----------------------------
rows, prev_df = [], None
files = list_files()
if not files:
    sys.exit("No <n>-<Hf|Zr>.txt files found in " + DATA_DIR)

for idx, (n, hc, path) in enumerate(files, start=1):
    df = read_txt(path)
    t = df["thickness"].to_numpy(dtype=float) * SCALE
    row = dict(label="%d-%s" % (n, hc), half_cycle_index=idx, cycle=n, half=hc,
               nodes=len(t), min=np.nanmin(t), max=np.nanmax(t), mean_nodal=np.nanmean(t),
               median=np.nanmedian(t), std=np.nanstd(t))
    row["nonuniformity_pct"] = (row["max"] - row["min"]) / (2 * row["mean_nodal"]) * 100 \
        if row["mean_nodal"] != 0 else np.nan

    # per-node increment, only if it is the same node set (same count and same coordinates)
    row["inc_node_min"] = row["inc_node_max"] = row["inc_node_std"] = np.nan
    if prev_df is not None and len(prev_df) == len(df):
        same = True
        for c in ("X", "Y", "Z"):
            if c in df and c in prev_df and df[c].notna().all() and prev_df[c].notna().all():
                same = same and np.allclose(df[c], prev_df[c], rtol=1e-6, atol=1e-15)
        if same:
            d = t - prev_df["thickness"].to_numpy(dtype=float) * SCALE
            row["inc_node_min"], row["inc_node_max"], row["inc_node_std"] = \
                np.nanmin(d), np.nanmax(d), np.nanstd(d)
    prev_df = df
    rows.append(row)

stats = pd.DataFrame(rows)

# ---------------------- optional area-weighted values ---------------------
areal_path = os.path.join(DATA_DIR, "areal_summary.csv")
if os.path.exists(areal_path):
    a = pd.read_csv(areal_path)
    a = a.rename(columns={"mean_nodal": "areal_mean_nodal"})
    keep = [c for c in a.columns if c not in ("source_file",)]
    stats = stats.merge(a[keep], on="label", how="left")

# ---------------------------- increments ----------------------------------
mean_cols = [c for c in ("mean_nodal", "mean_area_weighted", "mean_area_weighted_manual",
                         "mean_volume_weighted", "volume_over_area") if c in stats.columns]
for c in mean_cols:
    vals = stats[c].astype(float) * (SCALE if c != "mean_nodal" else 1.0)
    # mean_nodal is already scaled above; the areal columns come straight from the CSV
    stats[c] = vals
    prev = vals.shift(1).fillna(0.0)
    stats["inc_" + c] = vals - prev

if BASIS not in mean_cols:
    print("BASIS '%s' not available, using mean_nodal" % BASIS)
    BASIS = "mean_nodal"
INC = "inc_" + BASIS

# ------------------- node variation in every half-cycle file --------------
# total thickness in the file: how far the max / min node is from the mean
stats["var_max_pct"] = pct(stats["max"], stats["mean_nodal"])     # positive
stats["var_min_pct"] = pct(stats["min"], stats["mean_nodal"])     # negative
stats["range_pct"] = stats["var_max_pct"] - stats["var_min_pct"]  # (max - min) / mean * 100
stats["cv_pct"] = stats["std"] / stats["mean_nodal"] * 100
# thickness added in that half cycle (only where node sets match)
stats["inc_var_max_pct"] = pct(stats["inc_node_max"], stats["inc_mean_nodal"])
stats["inc_var_min_pct"] = pct(stats["inc_node_min"], stats["inc_mean_nodal"])
stats["inc_cv_pct"] = stats["inc_node_std"] / stats["inc_mean_nodal"] * 100

# ---------------------------- per-cycle table -----------------------------
cyc_rows, prev_vals = [], {c: 0.0 for c in mean_cols}
for n, g in stats.groupby("cycle"):
    r = {"cycle": n}
    for h in HALF_CYCLES:
        gh = g[g["half"] == h]
        r["added_%s" % h] = float(gh[INC].iloc[0]) if len(gh) else np.nan
        r["total_after_%s" % h] = float(gh[BASIS].iloc[0]) if len(gh) else np.nan
    end = g[g["half"] == H2]
    if end.empty:
        continue
    r["total_after_cycle"] = float(end[BASIS].iloc[0])
    r["growth_per_cycle"] = r["total_after_cycle"] - prev_vals[BASIS]
    r[RATIO_COL] = r["added_%s" % H1] / r["added_%s" % H2] if r["added_%s" % H2] else np.nan
    r["%s_share_pct" % H1] = r["added_%s" % H1] / r["growth_per_cycle"] * 100 \
        if r["growth_per_cycle"] else np.nan
    # the other mean definitions (area / volume weighted), if present
    for c in mean_cols:
        if c == BASIS:
            prev_vals[c] = r["total_after_cycle"]
            continue
        v = float(end[c].iloc[0])
        r["total_after_cycle_" + c] = v
        r["growth_per_cycle_" + c] = v - prev_vals[c]
        prev_vals[c] = v
    cyc_rows.append(r)
cyc = pd.DataFrame(cyc_rows)

# ------------------- error % wrt mean and mode of deposition --------------
ref_mask = (cyc["cycle"] > SKIP_CYCLES).to_numpy()
err = pd.DataFrame({"cycle": cyc["cycle"], "used_for_reference": ref_mask})
summary = []
for key, name in [("added_%s" % h, "added by %s half cycle" % h) for h in HALF_CYCLES] + \
                 [("growth_per_cycle", "growth per cycle (%s)" % " + ".join(HALF_CYCLES))]:
    v = cyc[key].to_numpy(dtype=float)
    ref = v[ref_mask]
    m_mean = np.nanmean(ref) if np.isfinite(ref).any() else np.nan
    m_mode, how = mode_of(ref)
    err[key] = v
    err[key + "_err_vs_mean_pct"] = pct(v, m_mean)
    err[key + "_err_vs_mode_pct"] = pct(v, m_mode)
    e_mean = err[key + "_err_vs_mean_pct"][ref_mask]
    e_mode = err[key + "_err_vs_mode_pct"][ref_mask]
    summary.append(dict(quantity=name, column=key, unit=UNIT, n_cycles=int(np.isfinite(ref).sum()),
                        mean=m_mean, mode=m_mode, mode_method=how,
                        median=np.nanmedian(ref), std=np.nanstd(ref),
                        min=np.nanmin(ref), max=np.nanmax(ref),
                        cv_pct=np.nanstd(ref) / m_mean * 100 if m_mean else np.nan,
                        err_vs_mean_min_pct=np.nanmin(e_mean), err_vs_mean_max_pct=np.nanmax(e_mean),
                        err_vs_mean_abs_avg_pct=np.nanmean(np.abs(e_mean)),
                        err_vs_mode_min_pct=np.nanmin(e_mode), err_vs_mode_max_pct=np.nanmax(e_mode),
                        err_vs_mode_abs_avg_pct=np.nanmean(np.abs(e_mode))))

# same error % for every half-cycle row (Hf rows vs Hf reference, Zr rows vs Zr reference)
for h in HALF_CYCLES:
    s = next(d for d in summary if d["column"] == "added_%s" % h)
    m = (stats["half"] == h).to_numpy()
    stats.loc[m, "inc_err_vs_mean_pct"] = pct(stats.loc[m, INC], s["mean"])
    stats.loc[m, "inc_err_vs_mode_pct"] = pct(stats.loc[m, INC], s["mode"])

rv = cyc[RATIO_COL].to_numpy(dtype=float)[ref_mask]
r_mode, r_how = mode_of(rv)
summary.append(dict(quantity="%s:%s deposition thickness ratio" % (H1, H2), column=RATIO_COL, unit="-",
                    n_cycles=int(np.isfinite(rv).sum()), mean=np.nanmean(rv), mode=r_mode,
                    mode_method=r_how, median=np.nanmedian(rv), std=np.nanstd(rv),
                    min=np.nanmin(rv), max=np.nanmax(rv),
                    cv_pct=np.nanstd(rv) / np.nanmean(rv) * 100))
for col, name in [("var_max_pct", "node max vs mean, total thickness (%)"),
                  ("var_min_pct", "node min vs mean, total thickness (%)"),
                  ("inc_var_max_pct", "node max vs mean, added thickness (%)"),
                  ("inc_var_min_pct", "node min vs mean, added thickness (%)")]:
    for h in HALF_CYCLES:
        v = stats.loc[stats["half"] == h, col].to_numpy(dtype=float)
        if np.isfinite(v).any():
            summary.append(dict(quantity="%s, %s files" % (name, h), column=col, unit="%",
                                n_cycles=int(np.isfinite(v).sum()), mean=np.nanmean(v),
                                median=np.nanmedian(v), std=np.nanstd(v),
                                min=np.nanmin(v), max=np.nanmax(v)))
summary = pd.DataFrame(summary)

node_var = stats[["label", "cycle", "half", "nodes", "mean_nodal", "min", "max", "std",
                  "var_max_pct", "var_min_pct", "range_pct", "cv_pct",
                  "inc_mean_nodal", "inc_node_min", "inc_node_max",
                  "inc_var_max_pct", "inc_var_min_pct", "inc_cv_pct"]]

# ------------------------------ write files -------------------------------
stats.to_csv(os.path.join(OUT_DIR, "thickness_stats.csv"), index=False)
cyc.to_csv(os.path.join(OUT_DIR, "cycle_stats.csv"), index=False)
err.to_csv(os.path.join(OUT_DIR, "deposition_error.csv"), index=False)

definitions = [
    ("Unit", "thickness in %s, percentages in %%" % UNIT),
    ("Basis", "%s (mean over nodes unless an area/volume weighted mean was chosen)" % BASIS),
    ("Reference cycles", "mean and mode use cycles > %d (SKIP_CYCLES)" % SKIP_CYCLES),
    ("added_<half>", "mean thickness after this half cycle - mean after the previous half cycle"),
    ("total_after_cycle", "mean total thickness after the %s half cycle of that cycle (%s)" % (H2, " + ".join(HALF_CYCLES))),
    ("growth_per_cycle", "total_after_cycle(n) - total_after_cycle(n-1)"),
    (RATIO_COL, "added_%s / added_%s in the same cycle" % (H1, H2)),
    ("var_max_pct / var_min_pct", "(max - mean) / mean * 100 and (min - mean) / mean * 100 over the nodes of one file"),
    ("inc_var_max_pct / inc_var_min_pct", "same, for the per-node thickness added in that half cycle (only where node sets match)"),
    ("err_vs_mean_pct", "(value - mean of value over reference cycles) / mean * 100"),
    ("err_vs_mode_pct", "(value - mode of value over reference cycles) / mode * 100"),
    ("mode", "most frequent value rounded to %d decimals; histogram peak if no value repeats" % MODE_DECIMALS),
]


def write_excel(path):
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        summary.to_excel(xw, sheet_name="Summary", index=False)
        r0 = len(summary) + 3
        pd.DataFrame(definitions, columns=["term", "definition"]).to_excel(
            xw, sheet_name="Summary", index=False, startrow=r0)
        cyc.to_excel(xw, sheet_name="Cycle_Stats", index=False)
        err.to_excel(xw, sheet_name="Deposition_Error", index=False)
        node_var.to_excel(xw, sheet_name="Node_Variation", index=False)
        stats.to_excel(xw, sheet_name="HalfCycle_Stats", index=False)

        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
        hdr_fill = PatternFill("solid", fgColor="DDE5F0")
        for ws in xw.book.worksheets:
            header_rows = {1, r0 + 1} if ws.title == "Summary" else {1}
            for hr in header_rows:
                for cell in ws[hr]:
                    if cell.value is not None:
                        cell.font = Font(bold=True)
                        cell.fill = hdr_fill
            if ws.title != "Summary":
                ws.freeze_panes = "B2"
            for col in ws.iter_cols(min_row=1, max_row=ws.max_row):
                width = max(len(str(c.value)) if c.value is not None else 0 for c in col[:60])
                ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(width + 2, 9), 60)
                for c in col[1:]:
                    if isinstance(c.value, float):
                        c.number_format = "0.000000"


xlsx_path = os.path.join(OUT_DIR, EXCEL_NAME)
try:
    import openpyxl  # noqa: F401
    try:
        write_excel(xlsx_path)
    except PermissionError:          # workbook still open in Excel
        xlsx_path = os.path.join(OUT_DIR, EXCEL_NAME.replace(".xlsx", time.strftime("_%H%M%S.xlsx")))
        write_excel(xlsx_path)
except ImportError:
    xlsx_path = "(no Excel file: openpyxl missing -> run  %s -m pip install openpyxl )" % sys.executable
    print("WARNING: openpyxl is not installed, so the Excel workbook was skipped. Install it with:\n"
          "    \"%s\" -m pip install openpyxl\n" % sys.executable)

pd.set_option("display.width", 220, "display.max_columns", 30)
print(stats[["label", "nodes", "mean_nodal", "inc_mean_nodal", "var_max_pct", "var_min_pct"]].to_string(index=False))
print()
print(cyc[["cycle", "added_" + H1, "added_" + H2, "growth_per_cycle", "total_after_cycle", RATIO_COL]].to_string(index=False))
print()
print(summary[["quantity", "mean", "mode", "std", "min", "max"]].to_string(index=False))
print("\nSaved %s, thickness_stats.csv, cycle_stats.csv, deposition_error.csv in %s"
      % (os.path.basename(xlsx_path), OUT_DIR))

# -------------------------------- plots -----------------------------------
colors = {h: c for h, c in zip(HALF_CYCLES, ("tab:blue", "tab:red", "tab:green", "tab:orange"))}
x = stats["half_cycle_index"].to_numpy()
labels = stats["label"].tolist()


def zoom_y(ax, *arrays):
    """Set the y range to the spread of the data (plus a small margin), not from zero."""
    vals = np.concatenate([np.asarray(a, dtype=float).ravel() for a in arrays])
    vals = vals[np.isfinite(vals)]
    if vals.size == 0:
        return
    lo, hi = vals.min(), vals.max()
    span = hi - lo
    if span == 0:
        span = abs(hi) * 0.02 if hi != 0 else 1.0
    ax.set_ylim(lo - Y_PAD * span, hi + Y_PAD * span)


def thin_ticks(ax, positions, texts, rotation=60):
    """Show only about MAX_XLABELS tick labels so the axis never gets crowded."""
    step = max(1, int(np.ceil(len(positions) / float(MAX_XLABELS))))
    ax.set_xticks(positions[::step])
    ax.set_xticklabels(texts[::step], rotation=rotation, ha="right" if rotation else "center")


def save(fig, name):
    fig.tight_layout()
    fig.savefig(os.path.join(OUT_DIR, name), dpi=200)


def range_bars(ax, xs, lo, hi, color, lw=1.2, alpha=0.8):
    """Vertical line from the minimum to the maximum at every x ('scatter line' of the spread)."""
    ax.vlines(xs, lo, hi, color=color, lw=lw, alpha=alpha)
    ax.plot(xs, lo, "_", color=color, ms=8, alpha=alpha)
    ax.plot(xs, hi, "_", color=color, ms=8, alpha=alpha)


def ref_lines(ax, d, fmt="%.4g"):
    u = "" if d["unit"] == "-" else " " + d["unit"]
    ax.axhline(d["mean"], color="k", ls="--", lw=1, label=("mean = " + fmt + "%s") % (d["mean"], u))
    ax.axhline(d["mode"], color="tab:green", ls=":", lw=1.5, label=("mode = " + fmt + "%s") % (d["mode"], u))


def shade_skipped(ax):
    if SKIP_CYCLES > 0:
        ax.axvspan(0.5, SKIP_CYCLES + 0.5, color="grey", alpha=0.12, label="not in mean / mode")


sm = {d["column"]: d for d in summary.to_dict("records")}
cx = cyc["cycle"].to_numpy()
ctxt = [str(int(c)) for c in cx]

# 1) thickness added in each half cycle (Hf and Zr in one plot), zoomed y axis
fig, ax = plt.subplots(figsize=(FIG_W, FIG_H))
ys = [stats[INC]]
for h in HALF_CYCLES:
    m = (stats["half"] == h).to_numpy()
    if BASIS == "mean_nodal" and stats["inc_node_min"][m].notna().any():
        range_bars(ax, x[m], stats["inc_node_min"][m], stats["inc_node_max"][m], colors[h], alpha=0.5)
        ys += [stats["inc_node_min"][m], stats["inc_node_max"][m]]
    ax.plot(x[m], stats[INC][m], "o-", color=colors[h], ms=6, zorder=3,
            label="added in %s half cycle (mean)" % h)
if BASIS == "mean_nodal" and "inc_mean_area_weighted" in stats and stats["inc_mean_area_weighted"].notna().any():
    ax.plot(x, stats["inc_mean_area_weighted"], "ks", mfc="none", ms=7, label="area-weighted mean")
    ys.append(stats["inc_mean_area_weighted"])
ax.set_xlabel("half cycle"); ax.set_ylabel("thickness added in that half cycle (%s)" % UNIT)
ax.set_title("Deposition per half cycle (bars = per-node min to max, where node sets match)")
thin_ticks(ax, x, labels)
zoom_y(ax, *ys)
ax.grid(alpha=0.3); ax.legend()
save(fig, "plot_increment_per_halfcycle.png")

# 2) thickness added by the Hf / by the Zr half cycle, separately, against cycle number
fig, axes = plt.subplots(1, len(HALF_CYCLES), figsize=(FIG_W, FIG_H), squeeze=False)
for j, h in enumerate(HALF_CYCLES):
    s = stats[stats["half"] == h]
    sx = s["cycle"].to_numpy()
    ax = axes[0][j]
    ys = [s[INC]]
    if BASIS == "mean_nodal" and s["inc_node_min"].notna().any():
        range_bars(ax, sx, s["inc_node_min"], s["inc_node_max"], colors[h], alpha=0.5)
        ys += [s["inc_node_min"], s["inc_node_max"]]
    ax.plot(sx, s[INC], "o-", color=colors[h], ms=6, label="mean added")
    ref_lines(ax, sm["added_%s" % h])
    shade_skipped(ax)
    ax.set_title("Thickness added by the %s half cycle" % h)
    ax.set_xlabel("cycle"); ax.set_ylabel("added thickness (%s)" % UNIT)
    thin_ticks(ax, sx, [str(int(c)) for c in sx], rotation=0)
    zoom_y(ax, *ys); ax.grid(alpha=0.3); ax.legend()
save(fig, "plot_added_Hf_Zr_separate.png")

# 3) growth per cycle (Hf + Zr)
if len(cyc) >= 1:
    fig, ax = plt.subplots(figsize=(FIG_W, FIG_H))
    ax.plot(cx, cyc["growth_per_cycle"], "o-", color="tab:purple", label="growth per cycle")
    ref_lines(ax, sm["growth_per_cycle"])
    shade_skipped(ax)
    ax.set_xlabel("cycle"); ax.set_ylabel("growth per cycle, %s (%s)" % (" + ".join(HALF_CYCLES), UNIT))
    ax.set_title("Growth per cycle")
    thin_ticks(ax, cx, ctxt, rotation=0)
    zoom_y(ax, cyc["growth_per_cycle"]); ax.grid(alpha=0.3); ax.legend()
    save(fig, "plot_growth_per_cycle.png")

# 4) Hf:Zr deposition thickness ratio
if len(cyc) >= 1:
    fig, ax = plt.subplots(figsize=(FIG_W, FIG_H))
    ax.plot(cx, cyc[RATIO_COL], "o-", color="tab:brown", label="%s:%s ratio (added %s / added %s)" % (H1, H2, H1, H2))
    ref_lines(ax, sm[RATIO_COL])
    shade_skipped(ax)
    ax.set_xlabel("cycle"); ax.set_ylabel("%s : %s deposition thickness ratio (-)" % (H1, H2))
    ax.set_title("%s:%s deposition thickness ratio per cycle" % (H1, H2))
    thin_ticks(ax, cx, ctxt, rotation=0)
    zoom_y(ax, cyc[RATIO_COL]); ax.grid(alpha=0.3); ax.legend()
    save(fig, "plot_HfZr_ratio.png")

# 5) % variation of node max / min wrt mean in every half-cycle file
has_inc = stats["inc_var_max_pct"].notna().any()
fig, axes = plt.subplots(2 if has_inc else 1, 1, figsize=(FIG_W, FIG_H * (1.4 if has_inc else 1)), squeeze=False)
panels = [("var_max_pct", "var_min_pct", "total thickness in the file")]
if has_inc:
    panels.append(("inc_var_max_pct", "inc_var_min_pct", "thickness added in that half cycle"))
for ax, (cmax, cmin, what) in zip(axes[:, 0], panels):
    for h in HALF_CYCLES:
        m = (stats["half"] == h).to_numpy()
        ax.plot(x[m], stats[cmax][m], "^-", color=colors[h], ms=6, label="%s: (max - mean)/mean" % h)
        ax.plot(x[m], stats[cmin][m], "v--", color=colors[h], ms=6, label="%s: (min - mean)/mean" % h)
    ax.axhline(0, color="k", lw=0.8)
    ax.set_xlabel("half cycle"); ax.set_ylabel("variation wrt mean (%)")
    ax.set_title("Node-to-node variation of the %s: max and min wrt mean" % what)
    thin_ticks(ax, x, labels)
    zoom_y(ax, stats[cmax], stats[cmin], [0.0]); ax.grid(alpha=0.3); ax.legend(ncol=2)
save(fig, "plot_node_variation.png")

# 6) error % of the deposited thickness wrt mean and mode
if len(cyc) >= 1:
    keys = ["added_%s" % h for h in HALF_CYCLES] + ["growth_per_cycle"]
    fig, axes = plt.subplots(len(keys), 1, figsize=(FIG_W, FIG_H * 1.6), squeeze=False, sharex=True)
    for ax, k in zip(axes[:, 0], keys):
        d = sm[k]
        ax.plot(cx, err[k + "_err_vs_mean_pct"], "o-", color="k", ms=5,
                label="vs mean (%.4g %s)" % (d["mean"], UNIT))
        ax.plot(cx, err[k + "_err_vs_mode_pct"], "s--", color="tab:green", ms=5, mfc="none",
                label="vs mode (%.4g %s)" % (d["mode"], UNIT))
        ax.axhline(0, color="grey", lw=0.8)
        shade_skipped(ax)
        ax.set_ylabel("error (%)")
        ax.set_title("Error of %s wrt mean and mode" % d["quantity"])
        zoom_y(ax, err[k + "_err_vs_mean_pct"], err[k + "_err_vs_mode_pct"], [0.0])
        ax.grid(alpha=0.3); ax.legend()
    axes[-1, 0].set_xlabel("cycle")
    thin_ticks(axes[-1, 0], cx, ctxt, rotation=0)
    save(fig, "plot_deposition_error.png")

print("Plots saved in", OUT_DIR)
print("Plots: plot_increment_per_halfcycle, plot_added_Hf_Zr_separate, plot_growth_per_cycle, "
      "plot_HfZr_ratio, plot_node_variation, plot_deposition_error (.png)")
if "--noshow" not in sys.argv:
    plt.show()