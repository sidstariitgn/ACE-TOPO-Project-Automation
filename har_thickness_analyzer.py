"""
har_thickness_analyzer.py  --  run with normal Python (NOT inside CFD-VIEW):
        python har_thickness_analyzer.py  [DATA_DIR]  [--noshow]

Reads the files written by hzo_har_xy_cut_thickness.py (CFD-VIEW)
    DATA_DIR/x-cut/<n>-Hf.txt, <n>-Zr.txt      (X-Slice through the hole axis)
    DATA_DIR/y-cut/<n>-Hf.txt, <n>-Zr.txt      (Y-Slice through the hole axis)
and writes everything to DATA_DIR/OUT_SUBDIR/
    x-cut/        tables + plots of the X cut
    y-cut/        tables + plots of the Y cut
    comparison/   X vs Y vs combined mean (tables + plots)
    har_stats.xlsx   all tables in one workbook

MEASUREMENT STATIONS ("nodes" used for every cycle)
    The cut points are not the same from file to file (the film surface moves, and slice
    points are created where the plane crosses the HZO mesh). To follow the SAME location
    through all cycles, the thickness is interpolated at fixed stations on the original hole:
        top       N_TOP_STATIONS per side, at offsets R .. R + top extent from the hole axis
        sidewall  N_DEPTH_STATIONS depths on the left and on the right wall
        bottom    N_BOTTOM_STATIONS offsets from -R to +R across the bottom
    For each station and each file the nearest HZO node (hzo_node from CFD-VIEW) is reported.

DEFINITIONS (per station)
    thickness after a half cycle     interpolated "Oxide"-"HZO" value
    added_Hf(n) = thk(n-Hf) - thk(n-1 -Zr)       (thk before cycle 1 = 0)
    added_Zr(n) = thk(n-Zr) - thk(n-Hf)
    Hf:Zr ratio per cycle            added_Hf(n) / added_Zr(n)  (nan if |added_Zr| < RATIO_MIN_ADDED)
    total Hf:Zr ratio                sum_n added_Hf(n) / sum_n added_Zr(n)   over ALL cycles
    combined (X+Y)                   thickness: mean of X and Y at the same station
                                     sidewall: mean of left and right wall of each cut, then of X and Y
                                     total ratio: sum Hf(mean of X,Y) / sum Zr(mean of X,Y)
    step coverage                    bottom mean / top mean, sidewall(depth) / top mean
    X vs Y difference %              (X - Y) / mean(X, Y) * 100

Requires: numpy, pandas, matplotlib, openpyxl
"""
import os
import re
import sys
import time
import warnings
import numpy as np
import pandas as pd
import matplotlib
if "--noshow" in sys.argv:
    matplotlib.use("Agg")
import matplotlib.pyplot as plt

warnings.filterwarnings("ignore", category=RuntimeWarning)   # empty means / 0-division -> nan

# ----------------------------- USER SETTINGS -----------------------------
DATA_DIR    = r"C:\Tests\Calibration_w_o_recomb\Stats\HAR_out"  # folder containing x-cut and y-cut
CUTS        = ["x-cut", "y-cut"]
HALF_CYCLES = ["Hf", "Zr"]            # order inside one cycle
SCALE       = 1e10                    # thickness: m -> A
UNIT        = "A"
LEN_SCALE   = 1e9                     # depth / position: m -> nm
LEN_UNIT    = "nm"

N_DEPTH_STATIONS  = 60                # stations down each sidewall
N_BOTTOM_STATIONS = 21                # stations across the bottom (-R .. +R)
N_TOP_STATIONS    = 10                # stations on the top surface, per side
TOP_EXTENT        = None              # m beyond the hole edge for top stations; None = auto
BOTTOM_R_FRACTION = 0.5               # bottom mean uses stations with |offset| < this * R
KEY_DEPTH_FRACTIONS = [0.1, 0.25, 0.5, 0.75, 0.9]   # depths reported in the summary tables

SKIP_CYCLES     = 0                   # first N cycles left out of per-cycle ratio mean / std
RATIO_MIN_ADDED = 1e-3                # (UNIT) per-cycle ratio = nan if |added Zr| is below this
N_PROFILE_CURVES = 6                  # cycles drawn as profiles (evenly spaced, after the Zr half)

OUT_SUBDIR  = "har_stats_results"
EXCEL_NAME  = "har_stats.xlsx"
FIG_W, FIG_H = 14, 8
# --------------------------------------------------------------------------

if len(sys.argv) > 1 and not sys.argv[1].startswith("--"):
    DATA_DIR = sys.argv[1]
OUT_DIR = os.path.join(DATA_DIR, OUT_SUBDIR)
H1, H2 = HALF_CYCLES[0], HALF_CYCLES[-1]
hc_rank = {h: i for i, h in enumerate(HALF_CYCLES)}
CUT_COLORS = {"x-cut": "tab:blue", "y-cut": "tab:red", "combined": "k"}


# ------------------------------- reading ----------------------------------
def list_files(folder):
    out = []
    for f in os.listdir(folder):
        m = re.match(r"^(\d+)-([A-Za-z]+)\.txt$", f)
        if m and m.group(2) in hc_rank:
            out.append((int(m.group(1)), m.group(2), os.path.join(folder, f)))
    out.sort(key=lambda t: (t[0], hc_rank[t[1]]))
    return out


def read_geometry(path):
    with open(path) as f:
        for line in f:
            if line.startswith("# geometry:"):
                return {k: (v if k == "cut_axis" else float(v))
                        for k, v in re.findall(r"(\w+)=(\S+)", line)}
            if not line.startswith("#"):
                break
    sys.exit("No '# geometry:' line in %s - rerun the updated CFD-VIEW script." % path)


def read_cut(cut):
    folder = os.path.join(DATA_DIR, cut)
    files = list_files(folder)
    if not files:
        sys.exit("No <n>-<Hf|Zr>.txt files in " + folder)
    data = []
    for n, hc, path in files:
        df = pd.read_csv(path, sep="\t", comment="#")
        df["thickness"] = df["thickness"].astype(float) * SCALE
        data.append(("%d-%s" % (n, hc), n, hc, df))
    return read_geometry(files[0][2]), data


raw, geom = {}, None
for cut in CUTS:
    g, d = read_cut(cut)
    geom = geom or g
    raw[cut] = d
R, D = geom["R"], geom["D"]
print("Hole: R = %.4g %s, D = %.4g %s" % (R * LEN_SCALE, LEN_UNIT, D * LEN_SCALE, LEN_UNIT))

# --------------------------- measurement stations --------------------------
if TOP_EXTENT is None:
    ext = []
    for cut in CUTS:
        df = raw[cut][-1][3]
        f = df[df["region"] == "field"]
        for sgn in (-1, 1):
            o = f["offset"][np.sign(f["offset"]) == sgn].abs()
            if len(o):
                ext.append(o.max() - R)
    TOP_EXTENT = min(ext) if ext else R
st = []
for sd, sgn in (("left", -1), ("right", 1)):
    for o in np.linspace(R * 1.02, R + TOP_EXTENT * 0.98, N_TOP_STATIONS):
        st.append(("top", sd, "offset", sgn * o))
dep = (np.arange(N_DEPTH_STATIONS) + 0.5) / N_DEPTH_STATIONS * D
for sd in ("left", "right"):
    for d in dep:
        st.append(("sidewall", sd, "depth", d))
for o in np.linspace(-R, R, N_BOTTOM_STATIONS):
    st.append(("bottom", "left" if o < 0 else "right", "offset", o))
stations = pd.DataFrame(st, columns=["region", "side", "coord_name", "coord"])
stations.insert(0, "station", ["%s_%s_%03d" % (r, s, i) for i, (r, s) in
                               enumerate(zip(stations["region"], stations["side"]))])
stations["coord_" + LEN_UNIT] = stations["coord"] * LEN_SCALE


def interp_file(df):
    """Thickness and nearest HZO node at every station for one file."""
    thk = np.full(len(stations), np.nan)
    node = np.full(len(stations), -1, dtype=int)
    for i, s in stations.iterrows():
        if s["region"] == "sidewall":
            p = df[(df["region"] == "sidewall") & (df["side"] == s["side"])]
            xcol = "depth"
        elif s["region"] == "top":
            p = df[(df["region"] == "field") & (np.sign(df["offset"]) == (1 if s["side"] == "right" else -1))]
            xcol = "offset"
        else:
            p = df[df["region"] == "bottom"]
            xcol = "offset"
        if len(p) < 1:
            continue
        g = p.groupby(xcol)["thickness"].mean()
        xs, ys = g.index.to_numpy(float), g.to_numpy(float)
        if len(xs) == 1:
            if abs(xs[0] - s["coord"]) < 0.02 * D:
                thk[i] = ys[0]
        elif xs.min() <= s["coord"] <= xs.max():
            thk[i] = np.interp(s["coord"], xs, ys)
        j = (p[xcol] - s["coord"]).abs().idxmin()
        node[i] = int(p.at[j, "hzo_node"]) if "hzo_node" in p else -1
    return thk, node


# ------------------------------ per cut ------------------------------------
def analyse(cut):
    labels, cyc_of, half_of, T, N = [], {}, {}, [], []
    for label, n, hc, df in raw[cut]:
        t, nd = interp_file(df)
        labels.append(label); cyc_of[label] = n; half_of[label] = hc
        T.append(t); N.append(nd)
    thk = pd.DataFrame(np.array(T).T, columns=labels)
    nodes = pd.DataFrame(np.array(N).T, columns=labels)
    # increments along the half-cycle sequence
    prev = np.zeros(len(stations))
    added = {}
    for lab in labels:
        cur = thk[lab].to_numpy()
        added[lab] = cur - prev
        prev = cur
    cycles = sorted(set(cyc_of.values()))
    addH = pd.DataFrame({n: added.get("%d-%s" % (n, H1), np.full(len(stations), np.nan)) for n in cycles})
    addZ = pd.DataFrame({n: added.get("%d-%s" % (n, H2), np.full(len(stations), np.nan)) for n in cycles})
    with np.errstate(divide="ignore", invalid="ignore"):
        ratio = addH / addZ.where(addZ.abs() >= RATIO_MIN_ADDED)
    ref = [c for c in cycles if c > SKIP_CYCLES]
    sumH, sumZ = addH.sum(axis=1, min_count=1), addZ.sum(axis=1, min_count=1)
    final_lab = labels[-1]
    last_df = raw[cut][-1][3]
    summ = stations.copy()
    summ["final_label"] = final_lab
    summ["final_thickness"] = thk[final_lab]
    summ["sum_added_" + H1] = sumH
    summ["sum_added_" + H2] = sumZ
    summ["total_ratio_%s_%s" % (H1, H2)] = sumH / sumZ
    summ["ratio_per_cycle_mean"] = ratio[ref].mean(axis=1)
    summ["ratio_per_cycle_std"] = ratio[ref].std(axis=1)
    summ["ratio_per_cycle_min"] = ratio[ref].min(axis=1)
    summ["ratio_per_cycle_max"] = ratio[ref].max(axis=1)
    summ["hzo_node_final"] = nodes[final_lab]
    xyz = last_df.set_index("hzo_node")[["X", "Y", "Z"]] if "hzo_node" in last_df else None
    for c in ("X", "Y", "Z"):
        summ[c + "_final"] = [xyz[c].get(k, np.nan) if xyz is not None and k >= 0 else np.nan
                              for k in nodes[final_lab]] if xyz is not None else np.nan
    for df_, name in ((thk, "thk"), (nodes, "nodes"), (addH, "addH"), (addZ, "addZ"), (ratio, "ratio")):
        df_.index = stations["station"]
    return dict(labels=labels, cyc_of=cyc_of, half_of=half_of, thk=thk, nodes=nodes,
                addH=addH, addZ=addZ, ratio=ratio, summ=summ, cycles=cycles)


res = {cut: analyse(cut) for cut in CUTS}


# ------------------------- region / step coverage --------------------------
def masks():
    m = {}
    m["top"] = (stations["region"] == "top").to_numpy()
    m["bottom"] = ((stations["region"] == "bottom") &
                   (stations["coord"].abs() < BOTTOM_R_FRACTION * R)).to_numpy()
    for sd in ("left", "right"):
        m["side_" + sd] = ((stations["region"] == "sidewall") & (stations["side"] == sd)).to_numpy()
    return m


M = masks()


def wall_at(values, frac):
    """Mean of left/right wall at depth frac*D (interpolated over the depth stations)."""
    out = {}
    for sd in ("left", "right"):
        v = np.asarray(values)[M["side_" + sd]]
        ok = np.isfinite(v)
        out[sd] = np.interp(frac * D, dep[ok], v[ok]) if ok.sum() > 1 else np.nan
    out["mean"] = np.nanmean([out["left"], out["right"]])
    return out


def per_label_stepcov(r):
    rows = []
    for lab in r["labels"]:
        v = r["thk"][lab].to_numpy()
        top, bot = np.nanmean(v[M["top"]]), np.nanmean(v[M["bottom"]])
        row = dict(label=lab, cycle=r["cyc_of"][lab], half=r["half_of"][lab], top=top, bottom=bot,
                   stepcov_bottom=bot / top if top else np.nan)
        for fr in KEY_DEPTH_FRACTIONS:
            w = wall_at(v, fr)
            row["side@%.3gD" % fr] = w["mean"]
            row["stepcov_side@%.3gD" % fr] = w["mean"] / top if top else np.nan
        rows.append(row)
    return pd.DataFrame(rows)


for cut in CUTS:
    res[cut]["stepcov"] = per_label_stepcov(res[cut])

# ---------------------------- combined X + Y -------------------------------
combo_st = stations.copy()
comp_rows = []
fx, fy = res[CUTS[0]], res[CUTS[-1]]
tX, tY = fx["summ"]["final_thickness"].to_numpy(), fy["summ"]["final_thickness"].to_numpy()
rc = "total_ratio_%s_%s" % (H1, H2)
rX, rY = fx["summ"][rc].to_numpy(), fy["summ"][rc].to_numpy()
hX, hY = fx["summ"]["sum_added_" + H1].to_numpy(), fy["summ"]["sum_added_" + H1].to_numpy()
zX, zY = fx["summ"]["sum_added_" + H2].to_numpy(), fy["summ"]["sum_added_" + H2].to_numpy()
combo_st["thk_" + CUTS[0]] = tX
combo_st["thk_" + CUTS[-1]] = tY
combo_st["thk_combined"] = np.nanmean([tX, tY], axis=0)
combo_st["thk_diff_X_minus_Y"] = tX - tY
with np.errstate(divide="ignore", invalid="ignore"):
    combo_st["thk_diff_pct"] = (tX - tY) / combo_st["thk_combined"] * 100
    combo_st["ratio_" + CUTS[0]] = rX
    combo_st["ratio_" + CUTS[-1]] = rY
    combo_st["ratio_combined"] = np.nanmean([hX, hY], axis=0) / np.nanmean([zX, zY], axis=0)
    combo_st["ratio_diff_pct"] = (rX - rY) / combo_st["ratio_combined"] * 100


def stats_of(v):
    v = np.asarray(v, float)
    v = v[np.isfinite(v)]
    if not v.size:
        return dict(n=0, mean=np.nan, std=np.nan, min=np.nan, max=np.nan)
    return dict(n=v.size, mean=v.mean(), std=v.std(), min=v.min(), max=v.max())


# comparison summary: one row per quantity
comp = []
quantities = [("top thickness", M["top"], "thk"), ("bottom thickness", M["bottom"], "thk"),
              ("sidewall thickness, left", M["side_left"], "thk"),
              ("sidewall thickness, right", M["side_right"], "thk"),
              ("total %s:%s ratio, top" % (H1, H2), M["top"], "ratio"),
              ("total %s:%s ratio, sidewall" % (H1, H2), M["side_left"] | M["side_right"], "ratio"),
              ("total %s:%s ratio, bottom" % (H1, H2), M["bottom"], "ratio")]
for name, m, kind in quantities:
    a = combo_st["%s_%s" % (kind, CUTS[0])].to_numpy()[m]
    b = combo_st["%s_%s" % (kind, CUTS[-1])].to_numpy()[m]
    c = combo_st["%s_combined" % kind].to_numpy()[m]
    dp = combo_st["%s_diff_pct" % kind].to_numpy()[m]
    sa, sb, sc_, sd = stats_of(a), stats_of(b), stats_of(c), stats_of(dp)
    comp.append({"quantity": name, "unit": UNIT if kind == "thk" else "-",
                 "mean_" + CUTS[0]: sa["mean"], "std_" + CUTS[0]: sa["std"],
                 "mean_" + CUTS[-1]: sb["mean"], "std_" + CUTS[-1]: sb["std"],
                 "mean_combined": sc_["mean"], "std_combined": sc_["std"],
                 "min_combined": sc_["min"], "max_combined": sc_["max"],
                 "mean_diff_pct_XvsY": sd["mean"],
                 "mean_abs_diff_pct_XvsY": np.nanmean(np.abs(dp)) if sd["n"] else np.nan,
                 "max_abs_diff_pct_XvsY": np.nanmax(np.abs(dp)) if sd["n"] else np.nan,
                 "n_stations": sc_["n"]})
# key depths and step coverage (final file)
for fr in KEY_DEPTH_FRACTIONS:
    wa, wb = wall_at(tX, fr), wall_at(tY, fr)
    cm = np.nanmean([wa["mean"], wb["mean"]])
    comp.append({"quantity": "sidewall thickness @ %.3g D (%.4g %s), L/R mean" % (fr, fr * D * LEN_SCALE, LEN_UNIT),
                 "unit": UNIT, "mean_" + CUTS[0]: wa["mean"], "mean_" + CUTS[-1]: wb["mean"],
                 "mean_combined": cm,
                 "mean_diff_pct_XvsY": (wa["mean"] - wb["mean"]) / cm * 100 if cm else np.nan,
                 "left_" + CUTS[0]: wa["left"], "right_" + CUTS[0]: wa["right"],
                 "left_" + CUTS[-1]: wb["left"], "right_" + CUTS[-1]: wb["right"]})
sx, sy = res[CUTS[0]]["stepcov"].iloc[-1], res[CUTS[-1]]["stepcov"].iloc[-1]
for key, name in [("stepcov_bottom", "step coverage bottom/top")] + \
                 [("stepcov_side@%.3gD" % fr, "step coverage sidewall@%.3gD/top" % fr) for fr in KEY_DEPTH_FRACTIONS]:
    a, b = sx[key], sy[key]
    cm = np.nanmean([a, b])
    comp.append({"quantity": name + " (final)", "unit": "-", "mean_" + CUTS[0]: a,
                 "mean_" + CUTS[-1]: b, "mean_combined": cm,
                 "mean_diff_pct_XvsY": (a - b) / cm * 100 if cm else np.nan})
# left / right asymmetry inside each cut
for cut in CUTS:
    t = res[cut]["summ"]["final_thickness"].to_numpy()
    L, Rr = t[M["side_left"]], t[M["side_right"]]
    with np.errstate(divide="ignore", invalid="ignore"):
        a = (L - Rr) / ((L + Rr) / 2) * 100
    s = stats_of(a)
    comp.append({"quantity": "%s: sidewall left-right asymmetry (%%)" % cut, "unit": "%",
                 "mean_combined": s["mean"], "min_combined": s["min"], "max_combined": s["max"],
                 "n_stations": s["n"]})
comp = pd.DataFrame(comp)

# combined step coverage vs half cycle
scX, scY = res[CUTS[0]]["stepcov"], res[CUTS[-1]]["stepcov"]
sc_comb = scX[["label", "cycle", "half"]].copy()
for c in scX.columns[3:]:
    sc_comb[c + "_" + CUTS[0]] = scX[c].values
    sc_comb[c + "_" + CUTS[-1]] = scY.set_index("label").reindex(scX["label"])[c].values
    sc_comb[c + "_combined"] = np.nanmean([sc_comb[c + "_" + CUTS[0]], sc_comb[c + "_" + CUTS[-1]]], axis=0)

# combined per-station per-cycle ratio (from combined added thickness)
cH = (res[CUTS[0]]["addH"] + res[CUTS[-1]]["addH"]) / 2
cZ = (res[CUTS[0]]["addZ"] + res[CUTS[-1]]["addZ"]) / 2
c_ratio = cH / cZ.where(cZ.abs() >= RATIO_MIN_ADDED)

# ------------------------------ write tables -------------------------------
for sub in CUTS + ["comparison"]:
    os.makedirs(os.path.join(OUT_DIR, sub), exist_ok=True)


def with_info(df):
    info = stations.set_index("station")[["region", "side", "coord_name", "coord_" + LEN_UNIT]]
    return info.join(df)


for cut in CUTS:
    r, od = res[cut], os.path.join(OUT_DIR, cut)
    r["summ"].to_csv(os.path.join(od, "station_summary.csv"), index=False)
    with_info(r["thk"]).to_csv(os.path.join(od, "station_thickness.csv"))
    with_info(r["nodes"]).to_csv(os.path.join(od, "station_hzo_node.csv"))
    with_info(r["addH"]).to_csv(os.path.join(od, "station_added_%s.csv" % H1))
    with_info(r["addZ"]).to_csv(os.path.join(od, "station_added_%s.csv" % H2))
    with_info(r["ratio"]).to_csv(os.path.join(od, "station_ratio_per_cycle.csv"))
    r["stepcov"].to_csv(os.path.join(od, "step_coverage.csv"), index=False)
oc = os.path.join(OUT_DIR, "comparison")
combo_st.to_csv(os.path.join(oc, "comparison_stations.csv"), index=False)
comp.to_csv(os.path.join(oc, "comparison_summary.csv"), index=False)
sc_comb.to_csv(os.path.join(oc, "step_coverage_combined.csv"), index=False)

definitions = [
    ("Units", "thickness in %s, depth / position in %s" % (UNIT, LEN_UNIT)),
    ("Hole", "R = %.5g %s, D = %.5g %s, top stations up to %.5g %s beyond the edge"
     % (R * LEN_SCALE, LEN_UNIT, D * LEN_SCALE, LEN_UNIT, TOP_EXTENT * LEN_SCALE, LEN_UNIT)),
    ("station", "fixed position on the original hole: top (offset), sidewall (depth, left/right), bottom (offset)"),
    ("thickness", "\"Oxide\"-\"HZO\" interpolated at the station from the cut points of that file"),
    ("hzo_node", "nearest node of the HZO surface object (from CFD-VIEW) to the station in that file"),
    ("added_%s(n)" % H1, "thk(n-%s) - thk(n-1-%s); before cycle 1 the thickness is 0" % (H1, H2)),
    ("added_%s(n)" % H2, "thk(n-%s) - thk(n-%s)" % (H2, H1)),
    ("ratio per cycle", "added_%s / added_%s, nan where |added_%s| < %g %s" % (H1, H2, H2, RATIO_MIN_ADDED, UNIT)),
    ("total ratio", "sum over all cycles of added_%s / sum of added_%s" % (H1, H2)),
    ("combined", "mean of %s and %s at the same station; total ratio from the mean sums" % (CUTS[0], CUTS[-1])),
    ("diff %", "(%s - %s) / combined * 100" % (CUTS[0], CUTS[-1])),
    ("bottom mean", "stations with |offset| < %.2g R" % BOTTOM_R_FRACTION),
    ("step coverage", "bottom / top and sidewall(depth, L/R mean) / top"),
    ("per-cycle ratio stats", "cycles > %d (SKIP_CYCLES)" % SKIP_CYCLES),
]


def write_excel(path):
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        comp.to_excel(xw, sheet_name="Comparison_Summary", index=False)
        pd.DataFrame(definitions, columns=["term", "definition"]).to_excel(
            xw, sheet_name="Comparison_Summary", index=False, startrow=len(comp) + 3)
        combo_st.to_excel(xw, sheet_name="Comparison_Stations", index=False)
        sc_comb.to_excel(xw, sheet_name="StepCoverage_Combined", index=False)
        for cut in CUTS:
            k = "X" if cut.startswith("x") else "Y"
            r = res[cut]
            r["summ"].to_excel(xw, sheet_name="%s_Station_Summary" % k, index=False)
            with_info(r["thk"]).to_excel(xw, sheet_name="%s_Thickness" % k)
            with_info(r["nodes"]).to_excel(xw, sheet_name="%s_HZO_Node" % k)
            with_info(r["addH"]).to_excel(xw, sheet_name="%s_Added_%s" % (k, H1))
            with_info(r["addZ"]).to_excel(xw, sheet_name="%s_Added_%s" % (k, H2))
            with_info(r["ratio"]).to_excel(xw, sheet_name="%s_Ratio_per_cycle" % k)
            r["stepcov"].to_excel(xw, sheet_name="%s_StepCoverage" % k, index=False)
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
        fill = PatternFill("solid", fgColor="DDE5F0")
        for ws in xw.book.worksheets:
            for cell in ws[1]:
                if cell.value is not None:
                    cell.font = Font(bold=True); cell.fill = fill
            ws.freeze_panes = "B2"
            for col in ws.iter_cols(min_row=1, max_row=min(ws.max_row, 60)):
                w = max(len(str(c.value)) if c.value is not None else 0 for c in col)
                ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(w + 2, 9), 50)


xlsx = os.path.join(OUT_DIR, EXCEL_NAME)
try:
    import openpyxl  # noqa: F401
    try:
        write_excel(xlsx)
    except PermissionError:
        xlsx = xlsx.replace(".xlsx", time.strftime("_%H%M%S.xlsx"))
        write_excel(xlsx)
except ImportError:
    print("WARNING: openpyxl missing -> no Excel file.  \"%s\" -m pip install openpyxl" % sys.executable)

pd.set_option("display.width", 200, "display.max_columns", 20)
print(comp[["quantity", "mean_" + CUTS[0], "mean_" + CUTS[-1], "mean_combined",
            "mean_diff_pct_XvsY"]].to_string(index=False))

# --------------------------------- plots -----------------------------------
L = LEN_SCALE


def save(fig, folder, name):
    fig.tight_layout()
    fig.savefig(os.path.join(OUT_DIR, folder, name), dpi=200)


def profile_labels(r):
    """N_PROFILE_CURVES evenly spaced labels ending with the H2 half cycle, last one included."""
    ends = [l for l in r["labels"] if r["half_of"][l] == H2] or r["labels"]
    idx = np.unique(np.linspace(0, len(ends) - 1, min(N_PROFILE_CURVES, len(ends))).round().astype(int))
    return [ends[i] for i in idx]


def sel(m, col):
    return stations["coord"].to_numpy()[m] * L, col[m]


for cut in CUTS:
    r = res[cut]
    labs = profile_labels(r)
    shades = plt.cm.viridis(np.linspace(0.15, 0.9, len(labs)))

    # 1) thickness vs depth (sidewall)
    fig, ax = plt.subplots(figsize=(FIG_W, FIG_H))
    for lab, c in zip(labs, shades):
        v = r["thk"][lab].to_numpy()
        for sd, ls in (("left", "-"), ("right", "--")):
            x, y = sel(M["side_" + sd], v)
            ax.plot(x, y, ls, color=c, label="%s, %s wall" % (lab, sd))
    ax.set_xlabel("depth below original top surface (%s)" % LEN_UNIT)
    ax.set_ylabel("deposited thickness (%s)" % UNIT)
    ax.set_title("%s: sidewall thickness vs depth" % cut)
    ax.grid(alpha=0.3); ax.legend(ncol=2, fontsize=8)
    save(fig, cut, "plot_thickness_vs_depth.png")

    # 2) bottom thickness across the radius
    fig, ax = plt.subplots(figsize=(FIG_W, FIG_H))
    for lab, c in zip(labs, shades):
        x, y = sel((stations["region"] == "bottom").to_numpy(), r["thk"][lab].to_numpy())
        ax.plot(x, y, "o-", color=c, ms=4, label=lab)
    ax.axvline(0, color="grey", lw=0.8)
    ax.set_xlabel("position from hole axis along the cut (%s)  [-R .. +R]" % LEN_UNIT)
    ax.set_ylabel("deposited thickness (%s)" % UNIT)
    ax.set_title("%s: bottom thickness across the diameter" % cut)
    ax.grid(alpha=0.3); ax.legend()
    save(fig, cut, "plot_bottom_thickness_vs_radius.png")

    # 3) top thickness
    fig, ax = plt.subplots(figsize=(FIG_W, FIG_H))
    for lab, c in zip(labs, shades):
        v = r["thk"][lab].to_numpy()
        for sd in ("left", "right"):
            m = M["top"] & (stations["side"] == sd).to_numpy()
            x, y = sel(m, v)
            ax.plot(x, y, "o-", color=c, ms=4, label=lab if sd == "left" else None)
    ax.set_xlabel("position from hole axis along the cut (%s)  [|x| > R]" % LEN_UNIT)
    ax.set_ylabel("deposited thickness (%s)" % UNIT)
    ax.set_title("%s: top-surface thickness" % cut)
    ax.grid(alpha=0.3); ax.legend()
    save(fig, cut, "plot_top_thickness.png")

    # 4) Hf:Zr ratio vs depth
    fig, ax = plt.subplots(figsize=(FIG_W, FIG_H))
    s = r["summ"]
    for sd, col in (("left", "tab:blue"), ("right", "tab:red")):
        m = M["side_" + sd]
        x = stations["coord"].to_numpy()[m] * L
        mu, sg = s["ratio_per_cycle_mean"].to_numpy()[m], s["ratio_per_cycle_std"].to_numpy()[m]
        ax.fill_between(x, mu - sg, mu + sg, color=col, alpha=0.15,
                        label="%s wall: per-cycle ratio mean +/- std" % sd)
        ax.plot(x, mu, ":", color=col)
        ax.plot(x, s[rc].to_numpy()[m], "o-", color=col, ms=4,
                label="%s wall: total ratio (sum %s / sum %s)" % (sd, H1, H2))
    ax.set_xlabel("depth below original top surface (%s)" % LEN_UNIT)
    ax.set_ylabel("%s : %s deposited thickness ratio (-)" % (H1, H2))
    ax.set_title("%s: %s:%s ratio vs depth" % (cut, H1, H2))
    ax.grid(alpha=0.3); ax.legend(fontsize=8)
    save(fig, cut, "plot_HfZr_ratio_vs_depth.png")

    # 5) total Hf:Zr ratio at every station (top / sidewall / bottom)
    fig, axes = plt.subplots(1, 3, figsize=(FIG_W * 1.3, FIG_H * 0.8))
    v = s[rc].to_numpy()
    for sd in ("left", "right"):
        m = M["top"] & (stations["side"] == sd).to_numpy()
        axes[0].plot(*sel(m, v), "o-", color="tab:green")
        axes[1].plot(*sel(M["side_" + sd], v), "o-", ms=3, label=sd + " wall")
    axes[2].plot(*sel((stations["region"] == "bottom").to_numpy(), v), "o-", color="tab:purple")
    for ax, t, xl in zip(axes, ("top", "sidewall", "bottom"),
                         ("position (%s)" % LEN_UNIT, "depth (%s)" % LEN_UNIT, "position (%s)" % LEN_UNIT)):
        ax.set_title("%s: total %s:%s ratio, %s" % (cut, H1, H2, t))
        ax.set_xlabel(xl); ax.set_ylabel("sum %s / sum %s (-)" % (H1, H2)); ax.grid(alpha=0.3)
    axes[1].legend()
    save(fig, cut, "plot_total_HfZr_ratio_per_station.png")

    # 6) step coverage vs cycle
    sc = r["stepcov"]
    sc = sc[sc["half"] == H2]
    fig, ax = plt.subplots(figsize=(FIG_W, FIG_H))
    ax.plot(sc["cycle"], sc["stepcov_bottom"], "o-", color="k", label="bottom / top")
    for fr, c in zip(KEY_DEPTH_FRACTIONS, plt.cm.plasma(np.linspace(0.1, 0.85, len(KEY_DEPTH_FRACTIONS)))):
        ax.plot(sc["cycle"], sc["stepcov_side@%.3gD" % fr], ".-", color=c, label="sidewall @ %.3g D / top" % fr)
    ax.set_xlabel("cycle"); ax.set_ylabel("step coverage (-)")
    ax.set_title("%s: step coverage after each %s half cycle" % (cut, H2))
    ax.grid(alpha=0.3); ax.legend()
    save(fig, cut, "plot_step_coverage.png")
    plt.close("all")

# ----------------------------- comparison plots ----------------------------
cols = {CUTS[0]: CUT_COLORS["x-cut"], CUTS[-1]: CUT_COLORS["y-cut"], "combined": "k"}
wall = {}
for cut in CUTS:
    t = res[cut]["summ"]["final_thickness"].to_numpy()
    wall[cut] = np.nanmean([t[M["side_left"]], t[M["side_right"]]], axis=0)
wall["combined"] = np.nanmean([wall[CUTS[0]], wall[CUTS[-1]]], axis=0)
final_lab = res[CUTS[0]]["labels"][-1]

fig, axes = plt.subplots(2, 2, figsize=(FIG_W * 1.3, FIG_H * 1.3))
ax = axes[0, 0]
for k in (CUTS[0], CUTS[-1], "combined"):
    ax.plot(dep * L, wall[k], "-" if k != "combined" else "--", color=cols[k], lw=2 if k == "combined" else 1.3,
            label=k + (" (L/R wall mean)" if k != "combined" else " mean of X and Y"))
ax.set_title("Sidewall thickness vs depth (%s)" % final_lab)
ax.set_xlabel("depth (%s)" % LEN_UNIT); ax.set_ylabel("thickness (%s)" % UNIT)
for ax, region, title in ((axes[0, 1], "bottom", "Bottom thickness across the diameter"),
                          (axes[1, 0], "top", "Top-surface thickness")):
    m = (stations["region"] == region).to_numpy()
    for k, col in ((CUTS[0], "thk_" + CUTS[0]), (CUTS[-1], "thk_" + CUTS[-1]), ("combined", "thk_combined")):
        v = combo_st[col].to_numpy()
        if region == "top":
            for sd in ("left", "right"):
                mm = m & (stations["side"] == sd).to_numpy()
                ax.plot(*sel(mm, v), "o-" if k != "combined" else "s--", ms=4, color=cols[k],
                        label=k if sd == "left" else None)
        else:
            ax.plot(*sel(m, v), "o-" if k != "combined" else "s--", ms=4, color=cols[k], label=k)
    ax.set_title("%s (%s)" % (title, final_lab))
    ax.set_xlabel("position from hole axis (%s)" % LEN_UNIT); ax.set_ylabel("thickness (%s)" % UNIT)
ax = axes[1, 1]
for k, col in ((CUTS[0], "ratio_" + CUTS[0]), (CUTS[-1], "ratio_" + CUTS[-1]), ("combined", "ratio_combined")):
    v = combo_st[col].to_numpy()
    ax.plot(dep * L, np.nanmean([v[M["side_left"]], v[M["side_right"]]], axis=0),
            "-" if k != "combined" else "--", color=cols[k], label=k)
ax.set_title("Total %s:%s ratio vs depth (L/R wall mean)" % (H1, H2))
ax.set_xlabel("depth (%s)" % LEN_UNIT); ax.set_ylabel("sum %s / sum %s (-)" % (H1, H2))
for a in axes.ravel():
    a.grid(alpha=0.3); a.legend(fontsize=8)
save(fig, "comparison", "plot_comparison_X_Y_combined.png")

# X - Y difference % along the profile
fig, axes = plt.subplots(1, 3, figsize=(FIG_W * 1.3, FIG_H * 0.8))
d = combo_st["thk_diff_pct"].to_numpy()
for sd in ("left", "right"):
    axes[0].plot(*sel(M["top"] & (stations["side"] == sd).to_numpy(), d), "o-", color="tab:green")
    axes[1].plot(*sel(M["side_" + sd], d), "o-", ms=3, label=sd + " wall")
axes[2].plot(*sel((stations["region"] == "bottom").to_numpy(), d), "o-", color="tab:purple")
for ax, t in zip(axes, ("top", "sidewall (vs depth)", "bottom")):
    ax.axhline(0, color="k", lw=0.8)
    ax.set_title("(%s - %s) / mean, %s" % (CUTS[0], CUTS[-1], t))
    ax.set_ylabel("thickness difference (%)"); ax.grid(alpha=0.3)
axes[1].legend()
save(fig, "comparison", "plot_X_minus_Y_difference.png")

# combined step coverage vs cycle
s2 = sc_comb[sc_comb["half"] == H2]
fig, ax = plt.subplots(figsize=(FIG_W, FIG_H))
for k, ls in ((CUTS[0], ":"), (CUTS[-1], "--"), ("combined", "-")):
    ax.plot(s2["cycle"], s2["stepcov_bottom_" + k], ls, color=cols[k], label="bottom/top, " + k)
    ax.plot(s2["cycle"], s2["stepcov_side@%.3gD_%s" % (KEY_DEPTH_FRACTIONS[len(KEY_DEPTH_FRACTIONS) // 2], k)],
            ls, color=cols[k], alpha=0.5,
            label="sidewall@%.3gD/top, %s" % (KEY_DEPTH_FRACTIONS[len(KEY_DEPTH_FRACTIONS) // 2], k))
ax.set_xlabel("cycle"); ax.set_ylabel("step coverage (-)")
ax.set_title("Step coverage: X cut, Y cut and combined"); ax.grid(alpha=0.3); ax.legend()
save(fig, "comparison", "plot_step_coverage_comparison.png")

print("\nResults in", OUT_DIR)
if "--noshow" not in sys.argv:
    plt.show()
